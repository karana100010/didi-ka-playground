"""M6: business plans: derivation from the registers, the design doc's dairy example, live formulas that Excel
computes to the same figures, range checks on the agent's changes, and the download."""
import io
import json

import pytest
from openpyxl import load_workbook

from conftest import ROOT, WS, login
from test_m4 import run_turn


def her_rows(didi):
    import sys
    sys.path.insert(0, str(ROOT / 'server'))
    from connectors.excel import ExcelConnector
    conn = ExcelConnector(lambda: (ROOT / 'data/mock/lokos_mock.xlsx').read_bytes(), 'test-plan')
    return {t: conn.rows(t, where={'didi_id': didi}) for t in ('dar', 'enterprise', 'savings')}


def test_dairy_plan_matches_the_design_example():
    from plans import builder, templates
    assumptions = templates.derive('dairy', her_rows('D-001'))
    values = {k: v['value'] for k, v in assumptions.items()}
    assert {k: values[k] for k in ('cows_now', 'milk_per_cow', 'milk_price', 'feed_per_cow', 'vet_per_cow', 'cow_price',
                                   'own_contribution', 'interest_rate', 'tenure_months')} == {
        'cows_now': 2, 'milk_per_cow': 8, 'milk_price': 38, 'feed_per_cow': 130, 'vet_per_cow': 300, 'cow_price': 70000,
        'own_contribution': 10000, 'interest_rate': 12, 'tenure_months': 24}
    assert assumptions['milk_per_cow']['source'] == 'Digital Aajeevika Register, Apr–Sep 2026 average'
    assert assumptions['cow_price']['source'].startswith('Mock JEEViKA unit cost')
    r = builder.compute('dairy', values)
    assert (r['profit_now'], r['profit_after'], r['loan_amount'], round(r['emi']), round(r['extra_after_emi']), r['payback_months']) == \
           (9840, 14760, 60000, 2824, 2096, 15)


@pytest.mark.parametrize('didi,kind', [('D-002', 'tailoring'), ('D-004', 'kirana'), ('D-006', 'dairy')])
def test_other_templates_build(didi, kind):
    from plans import builder, templates
    assumptions = templates.derive(kind, her_rows(didi))
    results = builder.compute(kind, {k: v['value'] for k, v in assumptions.items()})
    data = builder.build(kind, assumptions, results, {'name': 'Test', 'didi_id': didi, 'state': 'Jharkhand'}, {})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ['Summary', 'Assumptions', 'Cash flow', 'Loan', 'Break-even', 'Lakhpati path', 'Provenance']
    assert 'Bihar (JEEViKA) norms' in wb['Summary']['A5'].value
    if didi == 'D-002':  # Rekha: 3 machines, 66 garments a month in the register -> 22 per machine at ₹250
        assert (assumptions['machines_now']['value'], assumptions['garments_per_machine']['value'],
                assumptions['price_per_garment']['value']) == (3, 22, 250)
    if didi == 'D-006':  # Geeta has no enterprise yet: a new dairy starts from template defaults
        assert assumptions['cows_now']['value'] == 0 and results['profit_now'] == 0


def test_workbook_has_live_formulas_that_excel_computes_the_same(tmp_path):
    from plans import builder, templates
    assumptions = templates.derive('dairy', her_rows('D-001'))
    values = {k: v['value'] for k, v in assumptions.items()}
    results = builder.compute('dairy', values)
    data = builder.build('dairy', assumptions, results, {'name': 'Sunita Devi', 'didi_id': 'D-001', 'state': 'Bihar'}, {})
    wb = load_workbook(io.BytesIO(data))
    assert wb.calculation.fullCalcOnLoad is True
    assert wb['Summary']['B13'].value == '=emi' and 'cows_now' in wb.defined_names
    assert wb['Loan']['C2'].value == '=emi' and wb['Assumptions']['C4'].font.color.rgb.endswith('1F4FB8')  # inputs in blue
    assert any('MOCK' in str(c.value) for c in wb['Provenance']['A'])
    formulas = pytest.importorskip('formulas')
    path = tmp_path / 'plan.xlsx'
    path.write_bytes(data)
    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()

    def cell(sheet, ref):
        return next(v.value[0][0] for k, v in solution.items() if k.upper().endswith(f"{sheet.upper()}'!{ref}"))
    assert [round(float(cell('SUMMARY', f'B{r}'))) for r in (8, 9, 12, 13, 15, 16)] == [9840, 14760, 60000, 2824, 2096, 15]
    assert abs(float(cell('LOAN', 'F25'))) < 0.01  # the schedule pays the loan off


def test_changes_are_range_checked():
    from plans import service, templates
    assumptions = templates.derive('dairy', her_rows('D-001'))
    refused = service.apply_changes('dairy', assumptions, {'cows_to_add': 2, 'milk_price': 5000, 'goats': 3, 'feed_per_cow': 'lots'})
    assert assumptions['cows_to_add'] == {'value': 2, 'source': 'Her words, through the agent'}
    assert len(refused) == 3 and assumptions['milk_price']['value'] == 38


def test_plan_end_to_end_with_download(client, fake_llm):
    session = login(client, 'D-001')
    plan = {'needs_records': True, 'business_plan': 'dairy', 'plan_changes': {'cows_to_add': 1}, 'purpose': 'other_question',
            'items': [], 'scheme_topics': [], 'steps': ['Build the dairy plan']}

    fake_llm.script += [json.dumps(plan)]   # no writer call: the plan's answer is worded by code
    events, end = run_turn(client, session, 'मेरी डेयरी बढ़ाने का बिज़नेस प्लान बना दो')
    card = next(e['data'] for e in events if e['type'] == 'consent_request')
    assert card['purpose']['id'] == 'business_plan' and [i['table'] for i in card['items']] == ['dar', 'enterprise', 'savings']
    file = next(e['data'] for e in events if e['type'] == 'file')
    assert file['name'] == 'business-plan-dairy-D-001.xlsx' and file['results']['emi'] == 2824
    answer = end['runs'][0]['result']['answer']
    assert answer.startswith('सुनीता जी, आपके डेयरी का प्लान तैयार है। अभी हर महीने मुनाफ़ा ₹9,840 है; 1 गाय और जोड़ने के बाद ₹14,760 होगा।')
    assert '₹2,824' in answer and '₹2,096' in answer and '15 महीनों' in answer and 'नकली (mock)' in answer
    metrics = end['runs'][0]['result']['metrics']
    assert metrics['llm_calls'] == 1 and metrics['grounded_numbers'] == '7/7'

    response = client.get(f'/v1/files/{file["file_id"]}', headers=WS)
    assert response.status_code == 200 and response.headers['content-disposition'].endswith('business-plan-dairy-D-001.xlsx"')
    assert load_workbook(io.BytesIO(response.content))['Summary']['B8'].value == '=profit_now'
    assert client.get(f'/v1/files/{file["file_id"]}', headers={'X-Workspace': 'someone-else-0002'}).status_code == 404
