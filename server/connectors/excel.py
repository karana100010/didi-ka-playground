"""Excel workbooks (.xlsx): each worksheet is a table and its first non-empty row holds the column names.
Google Sheets are read through this too, exported as .xlsx."""
import io
import zipfile
from collections import OrderedDict

from openpyxl import load_workbook

from .base import MAX_ROWS, Connector, ConnectorError, clean

_CACHE = OrderedDict()  # cache_key -> parsed workbook; small, per warm function instance


def parse(data):
    """Bytes of an .xlsx file -> {table: {'columns': [...], 'rows': [dict, ...]}}."""
    if not data.startswith(b'PK'):
        raise ConnectorError('not_excel', 'This is not an Excel .xlsx file. Save it as .xlsx and upload it again.')
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        raise ConnectorError('not_excel', f'Could not read this workbook: {error}.') from error
    tables = {}
    try:
        for ws in wb.worksheets:
            header, index, rows = None, [], []
            for values in ws.iter_rows(values_only=True):
                if header is None:
                    if not any(v not in (None, '') for v in values):
                        continue
                    header = []
                    for i, name in enumerate(values):
                        name = str(name).strip() if name is not None else ''
                        if name and name not in header:
                            header.append(name)
                            index.append(i)
                    continue
                if not any(v not in (None, '') for v in values):
                    continue
                rows.append({name: clean(values[i]) if i < len(values) else None for name, i in zip(header, index)})
                if len(rows) > MAX_ROWS * 10:
                    raise ConnectorError('too_large', f'Sheet "{ws.title}" has more than {MAX_ROWS * 10} rows.')
            if header:
                tables[ws.title] = {'columns': header, 'rows': rows}
    finally:
        wb.close()
    if not tables:
        raise ConnectorError('empty', 'This workbook has no sheets with a header row.')
    return tables


class ExcelConnector(Connector):
    def __init__(self, load_bytes, cache_key):
        self._load, self._key = load_bytes, cache_key

    def _book(self):
        if self._key not in _CACHE:
            _CACHE[self._key] = parse(self._load())
            while len(_CACHE) > 8:
                _CACHE.popitem(last=False)
        _CACHE.move_to_end(self._key)
        return _CACHE[self._key]

    def tables(self):
        return {name: t['columns'] for name, t in self._book().items()}

    def rows(self, table, columns=None, where=None, limit=None):
        columns = self.check(table, columns, where)
        out = []
        for row in self._book()[table]['rows']:
            if where and any(str(row.get(k)) != str(v) for k, v in where.items()):
                continue
            out.append({c: row.get(c) for c in columns})
            if len(out) >= (limit or MAX_ROWS):
                break
        return out


def forget(cache_key):
    _CACHE.pop(cache_key, None)
