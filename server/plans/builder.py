"""Builds the business-plan workbook with openpyxl.

Inputs are blue cells on the Assumptions sheet, each a named cell (cows_now, milk_price, …). Every other
number is a live formula over those names, so the didi or a cadre can change an input in Excel and see the
plan update. The workbook is set to recalculate fully when opened (no LibreOffice on Vercel). compute()
runs the same formulas in Python for the spoken summary.
"""
import io
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from . import templates

BLUE = Font(color='1F4FB8', bold=True)
HINDI = 'Nirmala UI'
HEAD = PatternFill('solid', fgColor='1E3A5F')
WASH = PatternFill('solid', fgColor='FAF1DE')
MOCK_NOTE = {'hi': 'इस प्लान की लागतें नकली (mock) हैं, JEEViKA के असली मानक नहीं। असली मानक मिलने पर बदलें।',
             'en': 'Unit costs in this plan are MOCK figures, not real JEEViKA norms. Replace them when the real norms arrive.'}


def compute(enterprise, values):
    """Results from assumption values, with the same formulas the workbook uses."""
    from tools import calculate
    names = dict(values)
    for key, formula in templates.TEMPLATES[enterprise]['formulas'].items():
        names[key] = _calc(calculate, formula, names)
    for key, formula in templates.RESULTS:
        if key == 'emi':
            names[key] = emi(names['loan_amount'], names['interest_rate'], names['tenure_months'])
        elif key == 'payback_months':
            names[key] = math.ceil(names['investment'] / names['extra_profit']) if names['extra_profit'] > 0 else None
        else:
            names[key] = _calc(calculate, formula, names)
    return {k: v for k, v in names.items() if k not in values}


def _calc(calculate, formula, names):
    out = calculate(formula, {k: v for k, v in names.items() if isinstance(v, (int, float))})
    if 'error' in out:
        raise ValueError(f'{formula}: {out["error"]}')
    return out['result']


def emi(loan, rate_pct, months):
    if loan <= 0:
        return 0
    r = rate_pct / 1200
    return loan / months if r == 0 else loan * r / (1 - (1 + r) ** -months)


def excel(formula):
    """Calculator syntax -> Excel: function names in capitals."""
    return '=' + re.sub(r'\b(max|min|round)\(', lambda m: m.group(1).upper() + '(', formula)


def build(enterprise, assumptions, results, didi, provenance):
    """assumptions: {key: {'value', 'source'}}. Returns the .xlsx bytes."""
    spec = templates.TEMPLATES[enterprise]
    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    sheet = wb.create_sheet('Assumptions')
    cashflow = wb.create_sheet('Cash flow')
    loan = wb.create_sheet('Loan')
    breakeven = wb.create_sheet('Break-even')
    lakhpati = wb.create_sheet('Lakhpati path')
    prov = wb.create_sheet('Provenance')
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = f'Business plan: {spec["name"]["en"]} for {didi.get("name")}'

    # Assumptions and the calculated model, as named cells
    _header(sheet, ['इनपुट', 'Input', 'Value', 'Unit', 'Source'])
    row = 2
    for key, hi, en, unit, _, _ in templates.assumption_specs(enterprise):
        item = assumptions[key]
        sheet.append([hi, en, item['value'], unit, item['source']])
        sheet.cell(row=row, column=1).font = Font(name=HINDI)
        sheet.cell(row=row, column=3).font = BLUE
        _name(wb, key, 'Assumptions', row)
        row += 1
    row += 1
    sheet.cell(row=row, column=1, value='हिसाब (यहाँ कुछ न लिखें)').font = Font(name=HINDI, bold=True)
    sheet.cell(row=row, column=2, value='Calculated (do not type here)').font = Font(bold=True)
    row += 1
    calculated = list(spec['formulas'].items()) + [
        (key, formula or {'emi': 'IF(loan_amount>0, PMT(interest_rate/1200, tenure_months, -loan_amount), 0)',
                          'payback_months': 'IF(extra_profit>0, ROUNDUP(investment/extra_profit, 0), "never")'}[key])
        for key, formula in templates.RESULTS]
    labels = {'revenue_now': ('अभी की मासिक बिक्री', 'monthly revenue now'), 'cost_now': ('अभी का मासिक खर्च', 'monthly costs now'),
              'revenue_after': ('प्लान के बाद मासिक बिक्री', 'monthly revenue after the plan'),
              'cost_after': ('प्लान के बाद मासिक खर्च', 'monthly costs after the plan'), 'investment': ('कुल लागत', 'total investment'),
              'profit_now': ('अभी का मासिक मुनाफ़ा', 'monthly profit now'), 'profit_after': ('प्लान के बाद मासिक मुनाफ़ा', 'monthly profit after'),
              'extra_profit': ('हर महीने का अतिरिक्त मुनाफ़ा', 'extra profit a month'), 'loan_amount': ('कर्ज़', 'loan'),
              'emi': ('मासिक किस्त', 'monthly instalment'), 'extra_after_emi': ('किस्त के बाद अतिरिक्त नकद', 'extra cash after the instalment'),
              'payback_months': ('लागत कितने महीनों में वापस', 'months to pay back the investment'),
              'annual_income_after': ('प्लान के बाद सालाना आमदनी', 'household income a year after the plan'),
              'monthly_income_after': ('प्लान के बाद हर महीने की आमदनी', 'household income a month after the plan')}
    for key, formula in calculated:
        hi, en = labels[key]
        sheet.append([hi, en, excel(formula) if not formula.startswith(('IF(', 'PMT(')) else '=' + formula, '₹' if key != 'payback_months' else 'months'])
        sheet.cell(row=row, column=1).font = Font(name=HINDI)
        sheet.cell(row=row, column=3).number_format = '#,##0'
        _name(wb, key, 'Assumptions', row)
        row += 1
    _widths(sheet, [32, 40, 14, 14, 70])

    # Summary
    summary['A1'] = f'{spec["name"]["hi"]} का बिज़नेस प्लान · {didi.get("name_local") or didi.get("name")}'
    summary['A1'].font = Font(name=HINDI, bold=True, size=14)
    summary['A2'] = f'Business plan: {spec["name"]["en"]} · {didi.get("name")} ({didi.get("didi_id")}) · {didi.get("shg_name", "")}'
    summary['A3'] = MOCK_NOTE['hi']
    summary['A3'].font = Font(name=HINDI, color='94650A', bold=True)
    summary['A4'] = MOCK_NOTE['en']
    summary['A4'].font = Font(color='94650A')
    if (didi.get('state') or 'Bihar') != 'Bihar':
        summary['A5'] = f'Costs follow Bihar (JEEViKA) norms; {didi.get("state")} norms may differ.'
    rows = [('profit_now', 'अभी हर महीने का मुनाफ़ा (किस्त से पहले)', 'Monthly profit now, before instalment'),
            ('profit_after', 'प्लान के बाद हर महीने का मुनाफ़ा', 'Monthly profit after the plan, before instalment'),
            ('investment', 'कुल लागत', 'Total investment'), ('own_contribution', 'अपनी बचत से', 'From her savings'),
            ('loan_amount', 'कर्ज़', 'Loan'), ('emi', 'मासिक किस्त', 'Monthly instalment'),
            ('tenure_months', 'किस्तों के महीने', 'Instalment months'),
            ('extra_after_emi', 'किस्त के बाद हर महीने अतिरिक्त नकद', 'Extra cash each month after the instalment'),
            ('payback_months', 'लागत कितने महीनों में वापस', 'Months to pay back the investment'),
            ('annual_income_after', 'प्लान के बाद सालाना पारिवारिक आमदनी', 'Household income a year with the plan')]
    _header(summary, ['', 'Value'], row=7, labels=['सारांश / Summary', 'Value'])
    for n, (key, hi, en) in enumerate(rows, start=8):
        summary.cell(row=n, column=1, value=f'{hi} · {en}').font = Font(name=HINDI)
        summary.cell(row=n, column=2, value=f'={key}').number_format = '#,##0'
    status_row = 8 + len(rows) + 1
    summary.cell(row=status_row, column=1,
                 value='=IF(AND(annual_income_after>=100000, monthly_income_after>10000), "लखपति की आय-सीमा पार · Above the Lakhpati income line", '
                       '"लखपति की आय-सीमा से नीचे · Below the Lakhpati income line")').font = Font(name=HINDI, bold=True)
    _widths(summary, [72, 16])

    # Cash flow, 24 months
    _header(cashflow, ['Month', 'Revenue', 'Costs', 'Profit before instalment', 'Instalment', 'Net cash', 'Compared with now', 'Cumulative extra'])
    for m in range(1, 25):
        r = m + 1
        cashflow.append([m, '=revenue_after', '=cost_after', f'=B{r}-C{r}', f'=IF(A{r}<=tenure_months, emi, 0)', f'=D{r}-E{r}',
                         f'=F{r}-profit_now', f'=G{r}' if m == 1 else f'=H{r - 1}+G{r}'])
        for c in range(2, 9):
            cashflow.cell(row=r, column=c).number_format = '#,##0'
    _widths(cashflow, [8, 14, 14, 22, 14, 14, 18, 18])

    # Loan schedule
    _header(loan, ['Month', 'Opening balance', 'Instalment', 'Interest', 'Principal', 'Closing balance'])
    months = int(assumptions['tenure_months']['value']) if results['loan_amount'] > 0 else 0
    for m in range(1, months + 1):
        r = m + 1
        opening = '=loan_amount' if m == 1 else f'=F{r - 1}'
        loan.append([m, opening, '=emi', f'=B{r}*interest_rate/1200', f'=C{r}-D{r}', f'=B{r}-E{r}'])
        for c in range(2, 7):
            loan.cell(row=r, column=c).number_format = '#,##0'
    if not months:
        loan.append(['No loan needed: her savings cover the investment.'])
    _widths(loan, [8, 18, 14, 12, 12, 18])

    # Break-even and sensitivity (costs held constant)
    breakeven['A1'] = 'Months to pay back the investment'
    breakeven['B1'] = '=payback_months'
    breakeven['A3'] = 'Monthly profit after the plan if prices or volume change (costs held constant)'
    breakeven['A3'].font = Font(bold=True)
    changes = [-0.2, 0, 0.2]
    breakeven.append(['Volume change →'] + [f'{int(v * 100):+d}%' for v in changes])
    for p in ([-0.1, 0, 0.1] if spec['price'] else [0]):
        breakeven.append([f'Price {int(p * 100):+d}%' if spec['price'] else 'Price as now'] +
                         [f'=revenue_after*{1 + p}*{1 + v}-cost_after' for v in changes])
    for row_cells in breakeven.iter_rows(min_row=5, max_row=breakeven.max_row, min_col=2, max_col=4):
        for cell in row_cells:
            cell.number_format = '#,##0'
    _widths(breakeven, [44, 14, 14, 14])

    # Lakhpati path
    lakhpati['A1'] = 'लखपति की राह · Lakhpati path'
    lakhpati['A1'].font = Font(name=HINDI, bold=True, size=13)
    path = [('Household net income now (last 12 months)', '=household_income_now'),
            ('Household income a year with the plan, while repaying', '=annual_income_after'),
            ('Average a month with the plan', '=monthly_income_after'),
            ('Household income a year after the loan is repaid', '=household_income_now+extra_profit*12'),
            ('Lakhpati line: ₹1,00,000 a year and over ₹10,000 a month, for four business cycles', 100000)]
    for n, (label, formula) in enumerate(path, start=3):
        lakhpati.cell(row=n, column=1, value=label)
        lakhpati.cell(row=n, column=2, value=formula).number_format = '#,##0'
    _widths(lakhpati, [70, 16])

    # Provenance
    _header(prov, ['Input', 'Value', 'Source', 'Register', 'Consent', 'Read at'])
    for key, _, en, _, _, _ in templates.assumption_specs(enterprise):
        item = assumptions[key]
        ref = provenance.get(key) or {}
        prov.append([en, item['value'], item['source'], ref.get('register', ''), ref.get('consent_id', ''), ref.get('read_at', '')])
    prov.append([])
    prov.append([MOCK_NOTE['en']])
    prov.append(['Her register values were read once, with her consent, to build this plan. The workbook keeps the figures '
                 'she asked for; the playground does not keep the rows.'])
    _widths(prov, [36, 12, 70, 28, 18, 22])

    for ws in wb.worksheets:
        ws.freeze_panes = 'A2' if ws.title not in ('Summary', 'Break-even', 'Lakhpati path') else None
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _header(ws, columns, row=1, labels=None):
    for c, title in enumerate(labels or columns, start=1):
        cell = ws.cell(row=row, column=c, value=title)
        cell.font = Font(bold=True, color='FFFFFF', name=HINDI)
        cell.fill = HEAD
        cell.alignment = Alignment(vertical='center')


def _name(wb, key, sheet, row):
    wb.defined_names[key] = DefinedName(key, attr_text=f"'{sheet}'!$C${row}")


def _widths(ws, widths):
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = width
